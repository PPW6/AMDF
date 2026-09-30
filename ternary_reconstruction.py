import os
import json
from openpyxl import Workbook, load_workbook
import pandas as pd
def prcocess_re(mat, pum, prc2rum, mat2prc, prc2pum, bo):
    mat2prc_list, prc2pum_list, prc2rum_list = [], [], []
    if bo == 1:
        for prc in mat2prc:
            if mat[0] == prc[0]:
                mat2prc_list.append(prc[1])
        for prc in prc2pum:
            if pum[1] == prc[1]:
                prc2pum_list.append(prc[0])
        names_1 = [str(d) for d in mat2prc_list]
        names_2 = [str(d) for d in prc2pum_list]
        prc_set = set(names_1) & set(names_2)
        prc_list = [eval(one) for one in list(prc_set)]
        prc_list = sorted(prc_list, key=lambda x: x['pos'][0])
    elif bo == 2:
        for prc in mat2prc:
            if mat[0] == prc[0]:
                mat2prc_list.append(prc[1])
        for prc in prc2rum:
            if mat[1] == prc[0]:
                prc2rum_list.append(prc[0])
        names_1 = [str(d) for d in mat2prc_list]
        names_3 = [str(d) for d in prc2rum_list]
        prc_set = set(names_1) & set(names_3)
        prc_list = [eval(one) for one in list(prc_set)]
        prc_list = sorted(prc_list, key=lambda x: x['pos'][0])
    elif bo == 3:
        for prc in prc2pum:
            if pum[1] == prc[1]:
                prc2pum_list.append(prc[0])
        prc_list = prc2pum_list

    rum_list=['None'] * len(prc_list)
    if len(prc_list) != 0:
        processes = ', '.join([prc['name'] for prc in prc_list])
        for i in range(len(prc_list)):
            for rum in prc2rum:
                if prc_list[i] == rum[0]:
                    rum_list[i] = rum[1]['name']
                    break
                else:
                    rum_list[i] = 'None'
        if rum_list != ['None'] * len(prc_list):
            process_nums = ', '.join(rum_list)
        else:
            process_nums = 'None'
    else:
        processes = 'None'
        process_nums = 'None'
    return processes, process_nums
def relation_extration(dic_list, process_multi):
    if not dic_list['property'] or not process_multi:
        return []

    if isinstance(process_multi[0], dict):
        process_multi_id = [item['id'] for item in process_multi]
    elif isinstance(process_multi[0], int):
        process_multi_id = process_multi
    # process_one_id = process_one
    #将mat和property的正确匹配，判断句子中是一个工艺流程还是多个工艺流程对比，如果是一个则将工艺流程拼接。
    all_dict = []
    for name in dic_list['class_name']:
        if name != 'other' and name != 'process_multiple':
            target_item = dic_list[name+'_relation']
    for items in target_item:
        if items[0]['id'] not in process_multi_id:
            mat2pum, pro2pum, prc2rum, mat2prc, prc2pum = [],[],[],[],[]
            for item in items:
                if item['relation'] == 'mat2pum':
                    mat2pum.append((item['h'],item['t']))
                elif item['relation'] == 'pro2pum':
                    pro2pum.append((item['h'], item['t']))
                elif item['relation'] == 'prc2rum':
                    prc2rum.append((item['h'], item['t']))
                elif item['relation'] == 'mat2prc':
                    mat2prc.append((item['h'], item['t']))
                elif item['relation'] == 'prc2pum':
                    prc2pum.append((item['h'], item['t']))
            dict_one = {
                'DOI': dic_list['doi'],
                'alloy': 'None',
                'process': 'None',
                'process_num': 'None',
                'property': 'None',
                'property_num': 'None',
                'text': item['text']
            }
            found_mat = False
            if mat2pum:
                for mat in mat2pum:
                    for pum in pro2pum:
                        if mat[1] == pum[1]:
                            found_mat=True
                            processes, process_nums = prcocess_re(mat, pum, prc2rum, mat2prc, prc2pum, 1)
                            dict_one = {
                                'DOI': dic_list['doi'],
                                'alloy': mat[0]['name'],
                                'process': processes,
                                'process_num': process_nums,
                                'property': pum[0]['name'],
                                'property_num': pum[1]['name'],
                                'text': item['text']
                            }
                            all_dict.append(dict_one)

                if found_mat == False:
                    for mat in mat2prc:
                        for pum in prc2pum:
                            if mat[1] == pum[0]:
                                found_mat = True
                                processes, process_nums = prcocess_re(mat, pum, prc2rum, mat2prc, prc2pum, 1)
                                property = 'None'
                                for pro in pro2pum:
                                    if pro[1] == pum[1]:
                                        property = pro[0]['name']
                                dict_one = {
                                    'DOI': dic_list['doi'],
                                    'alloy': mat[0]['name'],
                                    'process': processes,
                                    'process_num': process_nums,
                                    'property': property,####pro的名字还没找到需要找一下pro的名字
                                    'property_num': pum[1]['name'],
                                    'text': item['text']
                                }
                                all_dict.append(dict_one)
            elif mat2prc and not mat2pum:
                processes, process_nums = [], []
                for mat in mat2prc:
                    prc2rum_unique = prc2rum.copy()  # Create a copy of prc2rum to modify
                    for prc in prc2rum_unique[:]:  # Iterate over a copy of prc2rum_unique to avoid issues while deleting
                        if mat[1] == prc[0]:
                            # Call the prcocess_re function and get process and process_num
                            process, process_num = prcocess_re(mat, [],  prc2rum_unique,mat2prc, prc2pum, 2)
                            processes.append(process)
                            process_nums.append(process_num)

                            # Remove the matched item from prc2rum_unique to prevent re-matching
                            prc2rum_unique.remove(prc)
                if len(processes) == 0:
                    processes, process_nums='None','None'
                else:
                    processes, process_nums=', '.join(processes), ', '.join(process_nums)
                dict_one = {
                    'DOI': dic_list['doi'],
                    'alloy': mat[0]['name'],
                    'process': processes,
                    'process_num': process_nums,
                    'property': 'None',
                    'property_num': 'None',
                    'text': item['text']
                }
                all_dict.append(dict_one)
            elif pro2pum and not mat2prc and not mat2pum:
                for pum in pro2pum:
                    processes, process_nums = prcocess_re([], pum, prc2rum, mat2prc, prc2pum,3)
                    dict_one = {
                        'DOI': dic_list['doi'],
                        'alloy': 'None',
                        'process': processes,
                        'process_num': process_nums,
                        'property': pum[0]['name'],
                        'property_num': pum[1]['name'],
                        'text': item['text']
                    }
                    all_dict.append(dict_one)
            elif prc2rum and not pro2pum and not mat2prc and not mat2pum:
                processes = ', '.join([prc[0]['name'] for prc in prc2rum])
                process_nums = ', '.join([prc[1]['name'] for prc in prc2rum])
                dict_one = {
                    'DOI': dic_list['doi'],
                    'alloy': 'None',
                    'process': processes,
                    'process_num': process_nums,
                    'property': 'None',
                    'property_num': 'None',
                    'text': item['text']
                }
                all_dict.append(dict_one)
            elif not pro2pum and not mat2prc and not prc2rum and not mat2pum:
                all_dict.append(dict_one)
    return all_dict

def save2xls(output_path,text_with_prc_relations):
    wb = Workbook()
    wb = load_workbook(output_path)
    sheet = wb.active
    for i, line in enumerate(text_with_prc_relations, start=1):
        sheet.cell(row=i, column=1, value=line)
    wb.save(output_path)

def save2txt(output_path,text_with_prc_relations):
    text_with_prc_relations = '\n'.join(text_with_prc_relations)
    with open(output_path,'w', encoding='utf-8') as f:
        f.write(text_with_prc_relations)
def save2json(output_path,relation):
    with open(output_path,'w', encoding='utf-8') as f:
        for item in relation:
            json_str=json.dumps(item,ensure_ascii=False)
            f.write('{}\n'.format(json_str))
def save2excel(xls_path,relation):
    try:
        wb = load_workbook(xls_path)
    except FileNotFoundError:
        wb = Workbook()
    if len(relation) == 0:
        return
    # Select the active sheet
    sheet = wb.active

    key_title = [key for key in relation[0]]
    # Write headers if the sheet is new
    if sheet.max_row == 1 and sheet.max_column == 1:
        sheet.append(key_title)

    # Calculate the starting row based on the offset
    start_row = sheet.max_row+1
    # Write each article to the sheet
    for i, item in enumerate(relation):
        for n in range(len(key_title)):
            sheet.cell(row=start_row + i, column=n+1, value=item[key_title[n]])

    # Save the workbook
    wb.save(xls_path)

if __name__ == '__main__':
    process_path = '../output_data/process'
    property_path = '../output_data/property'
    length = os.listdir(property_path)
    length = [len for len in length if 'RE' in len]
    output_excel='../output_data/data.xlsx'
    for i in range(0, len(length)):
    # for i in range(237,238):
        # with open(os.path.join(process_path, file_name), 'r', encoding='utf-8') as f:
        #     process = json.load(f)
        #     relation = relation_extration(process)
        #     output_path = os.path.join(process_path, process['doi'] + '_EX.json')
        #     if relation:
        #         save2json(output_path, relation)
        #         save2excel(output_excel, relation)
        file_name = length[i]
        with open(os.path.join(process_path, file_name[:-7]+'NER.json'), 'r', encoding='utf-8') as f:
            process = json.load(f)
        with open(os.path.join(property_path, file_name), 'r', encoding='utf-8') as f:
            property = json.load(f)
        print(f'{i}, resolve {file_name}')
        relation = relation_extration(property, process['process_multiple'])
        output_path = os.path.join(property_path, file_name[:-8]+'_EX.json')
        if relation:
            save2json(output_path, relation)


    re_length = os.listdir(property_path)
    re_length = [len for len in re_length if 'EX' in len]
    relation_ = []
    for i in range(0, len(re_length)):
        file_name = re_length[i]
        print(file_name)
        with open(os.path.join(property_path, file_name), 'r', encoding='utf-8') as f:
            lines = f.readlines()
            dict = [json.loads(line.strip()) for line in lines]
            relation_ += dict
    relation_ = [item for item in relation_ if item['property_num'] != 'None']
    save2excel(output_excel, relation_)
    # relation_mpa = []
    # for num in relation_:
    #     if num['alloy'] != 'None' and 'mpa' in num['property_num'].lower():  # 接下来用hobbs算法解决B-MAT和B-DSC的问题！！！！ 主要丢失MAT太多,在NER识别完之后就用hobbs算法。
    #         relation_mpa.append(num)
    #
    # save2excel(output_excel, relation_mpa)

